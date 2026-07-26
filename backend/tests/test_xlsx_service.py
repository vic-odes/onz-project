"""Tests du service d'export Excel du budget."""
import io

import openpyxl

from schemas.generated import GeneratedContent
from services import xlsx_service


def _generated(full: bool = True) -> dict:
    payload = {
        "introduction": "x",
        "cadre_logique": {},
        "parties_prenantes": [],
        "activites_detaillees": [],
        "chronogramme": [],
        "budget": {
            "lignes": [
                {"categorie": "Travaux", "description": "Forages", "montant_usd": 800000, "pourcentage": 80},
                {"categorie": "RH", "description": "Équipe", "montant_usd": 200000, "pourcentage": 20},
            ],
            "total_usd": 1000000,
            "couts_directs": 700000,
            "couts_indirects": 300000,
        },
        "analyse_cout_benefice": {},
        "risques": [],
        "communication": "x",
    }
    if full:
        payload["budget"]["par_annee"] = [
            {"annee": "Année 1", "montant_usd": 500000, "pourcentage": 50},
            {"annee": "Année 2", "montant_usd": 500000, "pourcentage": 50},
        ]
        payload["budget"]["par_partenaire"] = [
            {"partenaire": "ONG", "montant_usd": 900000, "pourcentage": 90},
        ]
        payload["plan_financement"] = {
            "lignes": [{"financeur": "AFD", "type": "Subvention", "montant_usd": 800000, "pourcentage": 80}],
            "total_usd": 1000000,
            "taux_cofinancement": "20 %",
        }
        payload["analyse_cout_benefice"] = {
            "van": 250000, "tri": 14.5, "ratio_cout_benefice": 1.8,
            "delai_retour_annees": 3.2, "benefices_annee1_usd": 120000,
        }
    return GeneratedContent.model_validate(payload).model_dump()


def _open(data: bytes):
    return openpyxl.load_workbook(io.BytesIO(data))


def test_full_workbook_sheets_and_values():
    data = xlsx_service.create_budget_workbook({"nom": "Mali"}, _generated(full=True))
    assert data[:2] == b"PK"
    wb = _open(data)
    assert set(wb.sheetnames) == {
        "Budget", "Budget par année", "Budget par partenaire",
        "Plan de financement", "Coût-bénéfice",
    }
    # La feuille Budget contient les catégories et un total
    ws = wb["Budget"]
    values = [v for row in ws.iter_rows(values_only=True) for v in row]
    assert "Travaux" in values
    assert "TOTAL" in values
    assert 1000000 in values


def test_minimal_workbook_only_budget_sheet():
    """Sans ventilations ni plan ni CBA, seule la feuille Budget est créée."""
    data = xlsx_service.create_budget_workbook({"nom": "Mini"}, _generated(full=False))
    wb = _open(data)
    assert wb.sheetnames == ["Budget"]


def test_cout_benefice_values():
    data = xlsx_service.create_budget_workbook({"nom": "Mali"}, _generated(full=True))
    wb = _open(data)
    values = [v for row in wb["Coût-bénéfice"].iter_rows(values_only=True) for v in row]
    assert 250000 in values   # VAN
    assert 14.5 in values     # TRI


def test_empty_generated_still_produces_valid_file():
    """Un contenu sans budget ne doit pas crasher (feuille Budget vide + total 0)."""
    data = xlsx_service.create_budget_workbook({"nom": "X"}, {"budget": {}})
    assert data[:2] == b"PK"
    wb = _open(data)
    assert "Budget" in wb.sheetnames
