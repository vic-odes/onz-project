from docx import Document
from docx.shared import Pt, RGBColor, Inches, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from datetime import datetime
import io


BLEU_MARINE = RGBColor(0x1B, 0x3A, 0x5C)
VERT_SAUGE = RGBColor(0x4A, 0x7C, 0x59)
GRIS_CLAIR = RGBColor(0xF0, 0xF0, 0xF0)

_ROMAN = [
    "", "I", "II", "III", "IV", "V", "VI", "VII", "VIII", "IX", "X",
    "XI", "XII", "XIII", "XIV", "XV", "XVI", "XVII", "XVIII", "XIX", "XX",
]


def _roman(n: int) -> str:
    """Chiffre romain pour la numérotation dynamique des sections (1..20)."""
    return _ROMAN[n] if 0 < n < len(_ROMAN) else str(n)


def _fmt_usd(value) -> str:
    try:
        return f"${float(value):,.0f} USD"
    except (TypeError, ValueError):
        return "N/A"


def _set_cell_bg(cell, hex_color: str):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    tcPr.append(shd)


def _add_header_footer(doc: Document, project_name: str):
    section = doc.sections[0]

    # En-tête
    header = section.header
    header.is_linked_to_previous = False
    p = header.paragraphs[0]
    p.text = f"ONZ Projet — {project_name}"
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = p.runs[0]
    run.font.size = Pt(9)
    run.font.color.rgb = BLEU_MARINE
    run.font.name = "Calibri"

    # Pied de page avec numérotation
    footer = section.footer
    footer.is_linked_to_previous = False
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    run.font.size = Pt(9)
    run.font.name = "Calibri"
    fldChar1 = OxmlElement("w:fldChar")
    fldChar1.set(qn("w:fldCharType"), "begin")
    instrText = OxmlElement("w:instrText")
    instrText.text = "PAGE"
    fldChar2 = OxmlElement("w:fldChar")
    fldChar2.set(qn("w:fldCharType"), "end")
    run._r.append(fldChar1)
    run._r.append(instrText)
    run._r.append(fldChar2)


def _add_title(doc: Document, text: str, level: int = 1):
    p = doc.add_paragraph()
    p.style = f"Heading {level}"
    run = p.add_run(text)
    run.font.name = "Calibri"
    run.font.size = Pt(14 if level == 1 else 12)
    run.font.color.rgb = BLEU_MARINE if level == 1 else VERT_SAUGE
    run.bold = True
    return p


def _add_paragraph(doc: Document, text: str):
    p = doc.add_paragraph(text)
    for run in p.runs:
        run.font.name = "Calibri"
        run.font.size = Pt(11)
    return p


def _add_bullets(doc: Document, items):
    """Ajoute une liste à puces mise en forme (pattern répété dans le document)."""
    for item in items:
        if item is None or str(item).strip() == "":
            continue
        p = doc.add_paragraph(str(item), style="List Bullet")
        for run in p.runs:
            run.font.name = "Calibri"
            run.font.size = Pt(11)


def _add_flow_step(doc: Document, label: str, content: str, *, last: bool = False):
    """Un maillon de chaîne causale (théorie du changement) suivi d'une flèche."""
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(label)
    run.font.name = "Calibri"
    run.font.size = Pt(11)
    run.bold = True
    run.font.color.rgb = BLEU_MARINE
    if content:
        run2 = p.add_run(f" : {content}")
        run2.font.name = "Calibri"
        run2.font.size = Pt(11)
    if not last:
        arrow = doc.add_paragraph("↓")
        arrow.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for r in arrow.runs:
            r.font.size = Pt(12)
            r.font.color.rgb = VERT_SAUGE


def _add_table_with_headers(doc: Document, headers: list[str], rows: list[list[str]]):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    # En-tête du tableau
    hdr_cells = table.rows[0].cells
    for i, header in enumerate(headers):
        hdr_cells[i].text = header
        _set_cell_bg(hdr_cells[i], "1B3A5C")
        run = hdr_cells[i].paragraphs[0].runs[0]
        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        run.font.bold = True
        run.font.name = "Calibri"
        run.font.size = Pt(10)

    # Lignes de données
    for row_data in rows:
        row_cells = table.add_row().cells
        for i, cell_text in enumerate(row_data):
            row_cells[i].text = str(cell_text)
            run = row_cells[i].paragraphs[0].runs[0] if row_cells[i].paragraphs[0].runs else row_cells[i].paragraphs[0].add_run(str(cell_text))
            run.font.name = "Calibri"
            run.font.size = Pt(10)

    doc.add_paragraph()
    return table


def create_word_document(project_data: dict, generated: dict) -> bytes:
    doc = Document()

    # Marges
    section = doc.sections[0]
    section.top_margin = Cm(2.5)
    section.bottom_margin = Cm(2.5)
    section.left_margin = Cm(3)
    section.right_margin = Cm(2.5)

    project_name = project_data.get("nom", "Projet")
    _add_header_footer(doc, project_name)

    # Compteur de sections : numérotation romaine dynamique (les sections
    # optionnelles ne créent pas de trou dans la numérotation).
    _counter = {"n": 0}

    def section_title(title: str):
        _counter["n"] += 1
        return _add_title(doc, f"{_roman(_counter['n'])}. {title}", 1)

    # ------------------------------------------------------------------ #
    # PAGE DE GARDE
    # ------------------------------------------------------------------ #
    doc.add_paragraph()
    doc.add_paragraph()

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("ONZ PROJET")
    run.font.name = "Calibri"
    run.font.size = Pt(14)
    run.font.color.rgb = VERT_SAUGE
    run.bold = True

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(project_name.upper())
    run.font.name = "Calibri"
    run.font.size = Pt(22)
    run.font.color.rgb = BLEU_MARINE
    run.bold = True

    # Type de dossier (montage vs demande de financement)
    dossier_label = (
        "Dossier de demande de financement"
        if project_data.get("type_dossier") == "financement"
        else "Dossier de montage de projet"
    )
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(dossier_label)
    run.font.name = "Calibri"
    run.font.size = Pt(13)
    run.font.color.rgb = VERT_SAUGE
    run.italic = True

    doc.add_paragraph()

    infos = [
        ("Pays / Zone d'intervention", project_data.get("pays", "")),
        ("Secteur", project_data.get("secteur", "")),
        ("Bailleur cible", project_data.get("bailleur", "")),
        ("Durée", f"{project_data.get('duree_mois', 'N/A')} mois"),
        ("Budget estimé", f"${project_data.get('budget_total', 'N/A'):,} USD" if project_data.get("budget_total") else "N/A"),
        ("Date de génération", datetime.now().strftime("%d/%m/%Y")),
    ]
    for label, value in infos:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(f"{label} : ")
        run.font.name = "Calibri"
        run.font.size = Pt(12)
        run.font.color.rgb = BLEU_MARINE
        run.bold = True
        run2 = p.add_run(str(value))
        run2.font.name = "Calibri"
        run2.font.size = Pt(12)

    doc.add_page_break()

    # ------------------------------------------------------------------ #
    # INTRODUCTION
    # ------------------------------------------------------------------ #
    section_title("Introduction")
    _add_paragraph(doc, generated.get("introduction", ""))
    doc.add_paragraph()

    # ------------------------------------------------------------------ #
    # ARBRE À PROBLÈMES / OBJECTIFS (optionnel)
    # ------------------------------------------------------------------ #
    arbre = generated.get("arbre_problemes") or {}
    if any(arbre.get(k) for k in ("probleme_central", "causes", "consequences", "objectif_central")):
        section_title("Arbre à problèmes et arbre à objectifs")

        _add_title(doc, "Arbre à problèmes", 2)
        if arbre.get("probleme_central"):
            p = doc.add_paragraph()
            run = p.add_run(f"Problème central : {arbre.get('probleme_central')}")
            run.font.name = "Calibri"
            run.font.size = Pt(11)
            run.bold = True
            run.font.color.rgb = BLEU_MARINE
        if arbre.get("causes"):
            _add_title(doc, "Causes (racines)", 2)
            _add_bullets(doc, arbre.get("causes", []))
        if arbre.get("consequences"):
            _add_title(doc, "Conséquences (effets)", 2)
            _add_bullets(doc, arbre.get("consequences", []))

        if any(arbre.get(k) for k in ("objectif_central", "moyens", "fins")):
            _add_title(doc, "Arbre à objectifs", 2)
            if arbre.get("objectif_central"):
                p = doc.add_paragraph()
                run = p.add_run(f"Objectif central : {arbre.get('objectif_central')}")
                run.font.name = "Calibri"
                run.font.size = Pt(11)
                run.bold = True
                run.font.color.rgb = VERT_SAUGE
            if arbre.get("moyens"):
                _add_title(doc, "Moyens", 2)
                _add_bullets(doc, arbre.get("moyens", []))
            if arbre.get("fins"):
                _add_title(doc, "Fins", 2)
                _add_bullets(doc, arbre.get("fins", []))
        doc.add_paragraph()

    # ------------------------------------------------------------------ #
    # THÉORIE DU CHANGEMENT (optionnel)
    # ------------------------------------------------------------------ #
    toc = generated.get("theorie_changement") or {}
    if any(toc.get(k) for k in ("narratif", "probleme", "impact", "activites", "resultats")):
        section_title("Théorie du changement")
        if toc.get("narratif"):
            _add_paragraph(doc, toc.get("narratif"))
            doc.add_paragraph()

        # Chaîne causale Problème → Impact
        _add_flow_step(doc, "Problème", toc.get("probleme", ""))
        _add_flow_step(doc, "Causes", " · ".join(toc.get("causes", [])))
        _add_flow_step(doc, "Activités", " · ".join(toc.get("activites", [])))
        _add_flow_step(doc, "Résultats", " · ".join(toc.get("resultats", [])))
        _add_flow_step(doc, "Effets", " · ".join(toc.get("effets", [])))
        _add_flow_step(doc, "Impact", toc.get("impact", ""), last=True)

        if toc.get("hypotheses"):
            doc.add_paragraph()
            _add_title(doc, "Hypothèses sous-jacentes", 2)
            _add_bullets(doc, toc.get("hypotheses", []))
        doc.add_paragraph()

    # ------------------------------------------------------------------ #
    # CADRE LOGIQUE
    # ------------------------------------------------------------------ #
    section_title("Cadre Logique")
    cadre = generated.get("cadre_logique", {})

    _add_title(doc, "Objectif Global", 2)
    _add_paragraph(doc, cadre.get("objectif_global", ""))

    _add_title(doc, "Objectifs Spécifiques", 2)
    _add_bullets(doc, cadre.get("objectifs_specifiques", []))

    _add_title(doc, "Résultats Attendus", 2)
    _add_bullets(doc, cadre.get("resultats", []))

    _add_title(doc, "Activités Principales", 2)
    _add_bullets(doc, cadre.get("activites", []))

    _add_title(doc, "Indicateurs SMART", 2)
    indicators = cadre.get("indicateurs_smart", [])
    if indicators:
        _add_table_with_headers(
            doc,
            ["Indicateur"],
            [[ind] for ind in indicators],
        )

    _add_title(doc, "Sources de Vérification", 2)
    _add_bullets(doc, cadre.get("sources_verification", []))

    _add_title(doc, "Hypothèses et Conditions Préalables", 2)
    _add_bullets(doc, cadre.get("hypotheses", []))

    # Matrice complète UE/AFD (Logique / Indicateur / Baseline / Cible / Source / Hypothèse)
    matrice = cadre.get("matrice", [])
    if matrice:
        _add_title(doc, "Matrice du Cadre Logique (format UE/AFD)", 2)
        rows = [
            [
                m.get("niveau", ""),
                m.get("logique_intervention", ""),
                m.get("indicateurs", ""),
                m.get("baseline", ""),
                m.get("cible", ""),
                m.get("sources_verification", ""),
                m.get("hypotheses", ""),
            ]
            for m in matrice
        ]
        _add_table_with_headers(
            doc,
            ["Niveau", "Logique d'intervention", "Indicateurs", "Baseline", "Cible", "Sources", "Hypothèses"],
            rows,
        )

    doc.add_page_break()

    # ------------------------------------------------------------------ #
    # PARTIES PRENANTES
    # ------------------------------------------------------------------ #
    section_title("Analyse des Parties Prenantes")
    parties = generated.get("parties_prenantes", [])
    if parties:
        rows = [
            [
                p.get("nom", ""),
                p.get("role", ""),
                p.get("interet", ""),
                p.get("influence", ""),
                p.get("quadrant", ""),
            ]
            for p in parties
        ]
        _add_table_with_headers(
            doc,
            ["Acteur", "Rôle", "Intérêt", "Influence", "Quadrant pouvoir/intérêt"],
            rows,
        )

    # ------------------------------------------------------------------ #
    # PLANIFICATION DES ACTIVITÉS
    # ------------------------------------------------------------------ #
    section_title("Planification des Activités")
    activites = generated.get("activites_detaillees", [])
    if activites:
        rows = [
            [
                a.get("titre", ""),
                a.get("description", ""),
                a.get("responsable", ""),
                a.get("duree", ""),
                a.get("objectif_lie", ""),
            ]
            for a in activites
        ]
        _add_table_with_headers(
            doc,
            ["Activité", "Description", "Responsable", "Durée", "Objectif lié"],
            rows,
        )

    _add_title(doc, "Chronogramme Trimestriel", 2)
    chrono = generated.get("chronogramme", [])
    if chrono:
        rows = [
            [c.get("trimestre", ""), " / ".join(c.get("activites", []))]
            for c in chrono
        ]
        _add_table_with_headers(doc, ["Trimestre", "Activités"], rows)

    doc.add_page_break()

    # ------------------------------------------------------------------ #
    # BUDGET
    # ------------------------------------------------------------------ #
    section_title("Budget Prévisionnel")
    budget = generated.get("budget", {})
    lignes = budget.get("lignes", [])
    if lignes:
        rows = [
            [
                l.get("categorie", ""),
                l.get("description", ""),
                f"${l.get('montant_usd', 0):,.0f}",
                f"{l.get('pourcentage', 0):.1f}%",
            ]
            for l in lignes
        ]
        _add_table_with_headers(
            doc,
            ["Catégorie", "Description", "Montant (USD)", "% du total"],
            rows,
        )

    p = doc.add_paragraph()
    run = p.add_run(
        f"Total budget : ${budget.get('total_usd', 0):,.0f} USD  |  "
        f"Coûts directs : ${budget.get('couts_directs', 0):,.0f}  |  "
        f"Coûts indirects : ${budget.get('couts_indirects', 0):,.0f}"
    )
    run.font.name = "Calibri"
    run.font.size = Pt(11)
    run.bold = True
    run.font.color.rgb = BLEU_MARINE

    doc.add_paragraph()

    # Ventilation annuelle (projet multi-annuel)
    par_annee = budget.get("par_annee", [])
    if par_annee:
        _add_title(doc, "Budget par année", 2)
        rows = [
            [
                b.get("annee", ""),
                f"${b.get('montant_usd', 0):,.0f}",
                f"{b.get('pourcentage', 0):.1f}%",
            ]
            for b in par_annee
        ]
        _add_table_with_headers(doc, ["Année", "Montant (USD)", "% du total"], rows)

    # Ventilation par partenaire
    par_partenaire = budget.get("par_partenaire", [])
    if par_partenaire:
        _add_title(doc, "Budget par partenaire", 2)
        rows = [
            [
                b.get("partenaire", ""),
                f"${b.get('montant_usd', 0):,.0f}",
                f"{b.get('pourcentage', 0):.1f}%",
            ]
            for b in par_partenaire
        ]
        _add_table_with_headers(doc, ["Partenaire", "Montant (USD)", "% du total"], rows)

    # ------------------------------------------------------------------ #
    # PLAN DE FINANCEMENT (optionnel)
    # ------------------------------------------------------------------ #
    plan = generated.get("plan_financement") or {}
    plan_lignes = plan.get("lignes", [])
    if plan_lignes:
        section_title("Plan de Financement")
        rows = [
            [
                f.get("financeur", ""),
                f.get("type", ""),
                f"${f.get('montant_usd', 0):,.0f}",
                f"{f.get('pourcentage', 0):.1f}%",
            ]
            for f in plan_lignes
        ]
        _add_table_with_headers(
            doc,
            ["Financeur", "Type", "Montant (USD)", "% du total"],
            rows,
        )
        p = doc.add_paragraph()
        run = p.add_run(
            f"Total financement : ${plan.get('total_usd', 0):,.0f} USD"
        )
        if plan.get("taux_cofinancement"):
            run2_text = f"  |  Taux de cofinancement : {plan.get('taux_cofinancement')}"
        else:
            run2_text = ""
        run.font.name = "Calibri"
        run.font.size = Pt(11)
        run.bold = True
        run.font.color.rgb = BLEU_MARINE
        if run2_text:
            run2 = p.add_run(run2_text)
            run2.font.name = "Calibri"
            run2.font.size = Pt(11)
            run2.bold = True
            run2.font.color.rgb = BLEU_MARINE
        doc.add_paragraph()

    doc.add_page_break()

    # ------------------------------------------------------------------ #
    # ANALYSE COÛT-BÉNÉFICE
    # ------------------------------------------------------------------ #
    section_title("Analyse Coût-Bénéfice")
    acb = generated.get("analyse_cout_benefice", {})
    _add_table_with_headers(
        doc,
        ["Indicateur", "Valeur"],
        [
            ["Valeur Actuelle Nette (VAN)", _fmt_usd(acb.get("van", 0))],
            ["Taux de Rentabilité Interne (TRI)", f"{acb.get('tri', 0):.1f}%"],
            ["Ratio Coût-Bénéfice", f"{acb.get('ratio_cout_benefice', 0):.2f}"],
            ["Délai de retour", f"{acb.get('delai_retour_annees', 0):.1f} an(s)"],
            ["Bénéfices estimés (année 1)", _fmt_usd(acb.get("benefices_annee1_usd", 0))],
            ["Scénario central", acb.get("scenario_central", "")],
            ["Scénario pessimiste", acb.get("scenario_pessimiste", "")],
        ],
    )
    if acb.get("analyse_sensibilite"):
        _add_title(doc, "Analyse de sensibilité", 2)
        _add_paragraph(doc, acb.get("analyse_sensibilite", ""))
    _add_title(doc, "Justification économique", 2)
    _add_paragraph(doc, acb.get("justification", ""))

    doc.add_page_break()

    # ------------------------------------------------------------------ #
    # GESTION DES RISQUES
    # ------------------------------------------------------------------ #
    section_title("Gestion des Risques")
    risques = generated.get("risques", [])
    if risques:
        rows = [
            [
                r.get("risque", ""),
                r.get("categorie", ""),
                r.get("probabilite", ""),
                r.get("impact", ""),
                r.get("niveau", ""),
                r.get("mitigation", ""),
            ]
            for r in risques
        ]
        _add_table_with_headers(
            doc,
            ["Risque", "Catégorie", "Probabilité", "Impact", "Niveau", "Mesure d'atténuation"],
            rows,
        )

    # ------------------------------------------------------------------ #
    # STRATÉGIE DE COMMUNICATION
    # ------------------------------------------------------------------ #
    section_title("Stratégie de Communication")
    _add_paragraph(doc, generated.get("communication", ""))

    # ------------------------------------------------------------------ #
    # PÉRENNISATION (optionnelle — flag inclure_perennisation)
    # ------------------------------------------------------------------ #
    perennisation = generated.get("perennisation", "")
    if perennisation and perennisation.strip():
        section_title("Pérennisation du Projet")
        _add_paragraph(doc, perennisation)

    # ------------------------------------------------------------------ #
    # NOTE CONCEPTUELLE (optionnelle)
    # ------------------------------------------------------------------ #
    note = generated.get("note_conceptuelle", "")
    if note and note.strip():
        doc.add_page_break()
        _add_title(doc, "Note Conceptuelle", 1)
        _add_paragraph(doc, note)

    # ------------------------------------------------------------------ #
    # RÉSUMÉ EXÉCUTIF (optionnel)
    # ------------------------------------------------------------------ #
    resume = generated.get("resume_executif", "")
    if resume and resume.strip():
        doc.add_page_break()
        _add_title(doc, "Résumé Exécutif", 1)
        _add_paragraph(doc, resume)

    # Sérialisation en bytes
    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf.read()
