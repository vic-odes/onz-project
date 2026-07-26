"""Génération de l'export Excel du budget (remarque #1 / #5 / #10).

Construit un classeur .xlsx à partir du contenu généré (`generated`) d'un projet :
- Budget par catégorie
- Budget par année (multi-annuel)
- Budget par partenaire
- Plan de financement
- Analyse coût-bénéfice

On ne crée que les feuilles dont les données existent ; une feuille « Budget »
est toujours présente pour ne jamais renvoyer un classeur vide.
"""

import io

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

BLEU_MARINE = "1B3A5C"
VERT_SAUGE = "4A7C59"
GRIS_CLAIR = "F0F0F0"

_HEADER_FONT = Font(name="Calibri", bold=True, color="FFFFFF", size=11)
_HEADER_FILL = PatternFill("solid", fgColor=BLEU_MARINE)
_TITLE_FONT = Font(name="Calibri", bold=True, color=BLEU_MARINE, size=14)
_TOTAL_FONT = Font(name="Calibri", bold=True, color=BLEU_MARINE, size=11)
_TOTAL_FILL = PatternFill("solid", fgColor=GRIS_CLAIR)
_CELL_FONT = Font(name="Calibri", size=11)
_THIN = Side(style="thin", color="D9D9D9")
_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)

_USD_FMT = '#,##0 "USD"'
_PCT_FMT = '0.0"%"'


def _f(value) -> float:
    try:
        return float(value or 0)
    except (TypeError, ValueError):
        return 0.0


def _style_header(ws, row: int, ncols: int):
    for c in range(1, ncols + 1):
        cell = ws.cell(row=row, column=c)
        cell.font = _HEADER_FONT
        cell.fill = _HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = _BORDER


def _autosize(ws, widths: list[int]):
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


def _write_table(ws, start_row: int, headers: list[str], rows: list[list], formats: list[str | None]) -> int:
    """Écrit un tableau (en-tête + lignes) et renvoie le numéro de la ligne suivante."""
    for j, h in enumerate(headers, 1):
        ws.cell(row=start_row, column=j, value=h)
    _style_header(ws, start_row, len(headers))

    r = start_row + 1
    for row_data in rows:
        for j, val in enumerate(row_data, 1):
            cell = ws.cell(row=r, column=j, value=val)
            cell.font = _CELL_FONT
            cell.border = _BORDER
            fmt = formats[j - 1] if j - 1 < len(formats) else None
            if fmt:
                cell.number_format = fmt
        r += 1
    return r


def _add_title(ws, row: int, text: str):
    cell = ws.cell(row=row, column=1, value=text)
    cell.font = _TITLE_FONT


def _sheet_categorie(wb: Workbook, budget: dict):
    ws = wb.active
    ws.title = "Budget"
    _add_title(ws, 1, "Budget prévisionnel par catégorie")
    lignes = budget.get("lignes", []) or []
    rows = [
        [l.get("categorie", ""), l.get("description", ""), _f(l.get("montant_usd")), _f(l.get("pourcentage"))]
        for l in lignes
    ]
    next_row = _write_table(
        ws, 3,
        ["Catégorie", "Description", "Montant", "% du total"],
        rows,
        [None, None, _USD_FMT, _PCT_FMT],
    )
    # Ligne de totaux
    total_cell = ws.cell(row=next_row, column=1, value="TOTAL")
    total_cell.font = _TOTAL_FONT
    total_cell.fill = _TOTAL_FILL
    ws.cell(row=next_row, column=2).fill = _TOTAL_FILL
    mt = ws.cell(row=next_row, column=3, value=_f(budget.get("total_usd")))
    mt.font = _TOTAL_FONT
    mt.fill = _TOTAL_FILL
    mt.number_format = _USD_FMT
    ws.cell(row=next_row, column=4).fill = _TOTAL_FILL

    # Coûts directs / indirects
    ws.cell(row=next_row + 2, column=1, value="Coûts directs").font = _CELL_FONT
    cd = ws.cell(row=next_row + 2, column=3, value=_f(budget.get("couts_directs")))
    cd.number_format = _USD_FMT
    cd.font = _CELL_FONT
    ws.cell(row=next_row + 3, column=1, value="Coûts indirects").font = _CELL_FONT
    ci = ws.cell(row=next_row + 3, column=3, value=_f(budget.get("couts_indirects")))
    ci.number_format = _USD_FMT
    ci.font = _CELL_FONT

    _autosize(ws, [26, 44, 16, 12])
    ws.freeze_panes = "A4"


def _simple_amount_sheet(wb: Workbook, title_sheet: str, title_text: str, label_header: str, items: list, label_key: str):
    ws = wb.create_sheet(title_sheet)
    _add_title(ws, 1, title_text)
    rows = [[it.get(label_key, ""), _f(it.get("montant_usd")), _f(it.get("pourcentage"))] for it in items]
    _write_table(ws, 3, [label_header, "Montant", "% du total"], rows, [None, _USD_FMT, _PCT_FMT])
    _autosize(ws, [32, 16, 12])
    ws.freeze_panes = "A4"


def _sheet_plan_financement(wb: Workbook, plan: dict):
    ws = wb.create_sheet("Plan de financement")
    _add_title(ws, 1, "Plan de financement")
    lignes = plan.get("lignes", []) or []
    rows = [
        [f.get("financeur", ""), f.get("type", ""), _f(f.get("montant_usd")), _f(f.get("pourcentage"))]
        for f in lignes
    ]
    next_row = _write_table(
        ws, 3,
        ["Financeur", "Type", "Montant", "% du total"],
        rows,
        [None, None, _USD_FMT, _PCT_FMT],
    )
    total_cell = ws.cell(row=next_row, column=1, value="TOTAL")
    total_cell.font = _TOTAL_FONT
    total_cell.fill = _TOTAL_FILL
    ws.cell(row=next_row, column=2).fill = _TOTAL_FILL
    mt = ws.cell(row=next_row, column=3, value=_f(plan.get("total_usd")))
    mt.font = _TOTAL_FONT
    mt.fill = _TOTAL_FILL
    mt.number_format = _USD_FMT
    ws.cell(row=next_row, column=4).fill = _TOTAL_FILL
    if plan.get("taux_cofinancement"):
        ws.cell(row=next_row + 2, column=1, value="Taux de cofinancement").font = _CELL_FONT
        ws.cell(row=next_row + 2, column=2, value=str(plan.get("taux_cofinancement"))).font = _CELL_FONT
    _autosize(ws, [30, 18, 16, 12])
    ws.freeze_panes = "A4"


def _sheet_cout_benefice(wb: Workbook, acb: dict):
    ws = wb.create_sheet("Coût-bénéfice")
    _add_title(ws, 1, "Analyse coût-bénéfice")
    rows = [
        ["Valeur Actuelle Nette (VAN)", _f(acb.get("van")), _USD_FMT],
        ["Taux de Rentabilité Interne (TRI)", _f(acb.get("tri")), _PCT_FMT],
        ["Ratio Coût-Bénéfice", _f(acb.get("ratio_cout_benefice")), "0.00"],
        ["Délai de retour (années)", _f(acb.get("delai_retour_annees")), "0.0"],
        ["Bénéfices estimés (année 1)", _f(acb.get("benefices_annee1_usd")), _USD_FMT],
    ]
    for j, h in enumerate(["Indicateur", "Valeur"], 1):
        ws.cell(row=3, column=j, value=h)
    _style_header(ws, 3, 2)
    r = 4
    for label, value, fmt in rows:
        lc = ws.cell(row=r, column=1, value=label)
        lc.font = _CELL_FONT
        lc.border = _BORDER
        vc = ws.cell(row=r, column=2, value=value)
        vc.font = _CELL_FONT
        vc.border = _BORDER
        vc.number_format = fmt
        r += 1
    _autosize(ws, [36, 20])
    ws.freeze_panes = "A4"


def create_budget_workbook(project_data: dict, generated: dict) -> bytes:
    """Construit le classeur Excel du budget et renvoie ses octets."""
    wb = Workbook()
    budget = generated.get("budget", {}) or {}

    _sheet_categorie(wb, budget)

    par_annee = budget.get("par_annee", []) or []
    if par_annee:
        _simple_amount_sheet(wb, "Budget par année", "Budget par année", "Année", par_annee, "annee")

    par_partenaire = budget.get("par_partenaire", []) or []
    if par_partenaire:
        _simple_amount_sheet(wb, "Budget par partenaire", "Budget par partenaire", "Partenaire", par_partenaire, "partenaire")

    plan = generated.get("plan_financement", {}) or {}
    if plan.get("lignes"):
        _sheet_plan_financement(wb, plan)

    acb = generated.get("analyse_cout_benefice", {}) or {}
    if any(_f(acb.get(k)) for k in ("van", "tri", "ratio_cout_benefice", "delai_retour_annees", "benefices_annee1_usd")):
        _sheet_cout_benefice(wb, acb)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.read()
